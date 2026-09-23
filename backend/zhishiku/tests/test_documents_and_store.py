from __future__ import annotations

# pytest 会自动收集名称以 test_ 开头的函数；正常启动 FastAPI 不会执行本文件。
# 本文件测试解析与存储边界，数据都在内存构造，不读取正式知识库。
# 阅读测试可按 AAA：Arrange（准备）→ Act（执行）→ Assert（断言）。

import io

import pytest
from docx import Document as DocxDocument

from backend.zhishiku.document_loader import (
    DocumentParseError,
    content_hash,
    parse_document,
    split_document,
)
from backend.zhishiku.vector_stories import KnowledgeStore


def make_docx() -> bytes:
    # 测试辅助函数在内存创建 DOCX 并返回 bytes，不依赖磁盘样本文件。
    stream = io.BytesIO()
    document = DocxDocument()
    document.add_heading("日本站品牌备案", level=1)
    document.add_paragraph("备案前应核对商标状态和站点要求。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "材料"
    table.cell(0, 1).text = "商标注册信息"
    document.add_paragraph("提交后在后台查看审核结果。")
    document.save(stream)
    return stream.getvalue()


def make_text_pdf(text: str = "Amazon VAT guide") -> bytes:
    """构造一个最小、带 Helvetica 文本层的 PDF，无需额外生成依赖。"""

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
    ]
    stream = ("BT /F1 12 Tf 30 100 Td (" + text + ") Tj ET").encode()
    objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(str(index).encode() + b" 0 obj\n" + obj + b"\nendobj\n")
    xref = len(data)
    data.extend(b"xref\n0 6\n0000000000 65535 f \n")
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend(b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n")
    data.extend(str(xref).encode() + b"\n%%EOF")
    return bytes(data)


def test_docx_preserves_heading_table_and_paragraph_order() -> None:
    # Arrange + Act：构造输入并调用正式解析入口。
    parsed = parse_document(make_docx(), "001_日本站品牌备案.docx")
    # Assert：不仅验证有文字，还验证段落、表格和后续段落的先后顺序。
    combined = "\n".join(body for _, body in parsed.sections)
    assert parsed.title == "日本站品牌备案"
    assert combined.index("备案前") < combined.index("材料 | 商标注册信息") < combined.index("提交后")
    assert "日本站品牌备案" in parsed.text


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "gb18030"])
def test_txt_supported_encodings(encoding: str) -> None:
    # parametrize 会分别传入三种编码，等价于生成三条测试但不复制函数。
    parsed = parse_document("欧洲站 VAT 注册要求".encode(encoding), "VAT.txt")
    assert "VAT" in parsed.text


def test_pdf_with_text_layer_and_scanned_pdf_error() -> None:
    assert "Amazon VAT guide" in parse_document(make_text_pdf(), "guide.pdf").text
    from pypdf import PdfWriter

    output = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(output)
    with pytest.raises(DocumentParseError, match="不执行 OCR"):
        parse_document(output.getvalue(), "scan.pdf")


def test_empty_corrupt_and_unsupported_files() -> None:
    with pytest.raises(DocumentParseError, match="为空"):
        parse_document(b"", "empty.txt")
    with pytest.raises(DocumentParseError, match="损坏"):
        parse_document(b"not-a-zip", "broken.docx")
    with pytest.raises(DocumentParseError, match="仅支持"):
        parse_document(b"abc", "data.csv")


def test_normalized_hash_and_chunk_ids_are_deterministic() -> None:
    first = parse_document("标题\r\n正文  内容".encode(), "a.txt")
    second = parse_document("标题\n正文 内容".encode(), "renamed.txt")
    assert content_hash(first.text) == content_hash(second.text)
    assert [item.chunk_id for item in split_document(first)] == [
        item.chunk_id for item in split_document(second)
    ]


class FakeCollection:
    # Fake 只实现 KnowledgeStore 真正用到的接口，不连接正式 Chroma。
    def __init__(self, backend) -> None:
        self.backend = backend

    def count(self) -> int:
        return len(self.backend.documents)


class FakeChroma:
    # documents 与 ids 只存在本测试内存中；add_documents 不会调用 Embedding。
    def __init__(self) -> None:
        self.documents = []
        self.ids = []
        self._collection = FakeCollection(self)

    def get(self, where=None, limit=None, include=None):
        matches = list(zip(self.ids, self.documents))
        if where and "content_hash" in where:
            matches = [
                pair for pair in matches
                if pair[1].metadata.get("content_hash") == where["content_hash"]
            ]
        if limit:
            matches = matches[:limit]
        return {
            "ids": [item[0] for item in matches],
            "metadatas": [item[1].metadata for item in matches],
        }

    def add_documents(self, documents, ids):
        self.documents.extend(documents)
        self.ids.extend(ids)
        return ids


def test_chroma_is_the_deduplication_source() -> None:
    # Arrange：把内存 Fake 注入延迟客户端位置。
    backend = FakeChroma()
    store = KnowledgeStore()
    store._store = backend
    parsed = parse_document("同一份亚马逊正文".encode(), "first.txt")

    # Act：相同正文连续写入两次；第二次故意换分类也应判重。
    first = store.index_document(parsed, "运营")
    second = store.index_document(parsed, "另一个分类")

    # Assert：验证首次入库、再次去重以及安全 metadata。
    assert first["status"] == "indexed"
    assert second["status"] == "duplicate"
    assert len(backend.ids) == first["chunk_count"]
    assert all(item.metadata["source"] == "first.txt" for item in backend.documents)
