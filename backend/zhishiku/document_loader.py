"""DOCX、PDF、TXT 的安全解析、清洗与切块。

这里把“文件字节 -> 纯文本 -> 知识块”集中处理。HTTP 上传和目录批量导入
共用同一实现，所以两条入库路径不会产生不同的去重结果或切块格式。
"""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document as DocxDocument
from docx.document import Document as DocxDocumentType
from docx.table import Table
from docx.text.paragraph import Paragraph
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

from .config_data import (
    CHUNK_OVERLAP,
    CHUNK_SIZE,
    MAX_DOCX_UNCOMPRESSED_BYTES,
    MAX_UPLOAD_BYTES,
)


SUPPORTED_SUFFIXES = {".docx", ".pdf", ".txt"}
# 分隔符按优先级从“大语义边界”排到“字符兜底”：先找段落、句子，
# 实在找不到边界才按空格或字符切。扩展新格式时也可复用这个顺序。
SPLIT_SEPARATORS = ["\n\n", "\n", "。", "；", "！", "？", ".", " ", ""]
_DATE_PATTERN = re.compile(r"(?<!\d)(20\d{2})[年./-](\d{1,2})[月./-](\d{1,2})日?")
_HEADING_PATTERN = re.compile(
    r"^(?:第[一二三四五六七八九十百0-9]+[章节部分篇]|"
    r"[一二三四五六七八九十]+[、.]|\d+(?:\.\d+)*[、.：:]?)\s*"
)


class DocumentParseError(ValueError):
    """文件可读取，但没有可用于知识库的正文或格式已经损坏。"""


# ParsedDocument 主要承载数据。@dataclass 会自动生成 __init__；slots=True
# 能阻止随意增加或拼错字段，也可降低大量文档对象占用的内存。
@dataclass(slots=True)
class ParsedDocument:
    """解析后的统一文档。

    text 是用于 SHA-256 去重的规范化全文；sections 保留章节名称，使切块
    后仍能把标题写回每块前面；published_at 无法识别时为空字符串，不使用
    文件修改时间冒充正文发布日期。
    """

    filename: str
    title: str
    text: str
    sections: list[tuple[str, str]]
    published_at: str


@dataclass(slots=True)
class KnowledgeChunk:
    """一个准备写入 Chroma 的知识块及其确定性 ID。"""

    chunk_id: str
    text: str
    chunk_index: int
    section: str


def normalize_text(text: str) -> str:
    """统一换行、空白和不可见字符，得到稳定的去重正文。"""

    # 所有格式先提取文字，再经过这里进入哈希和切块。统一清洗规则，才能让
    # 相同正文稳定地产生相同指纹；新增文件格式时也应该复用这个数据边界。
    text = text.replace("\ufeff", "").replace("\u200b", "").replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[\t\u00a0\u3000 ]+", " ", line).strip() for line in text.split("\n")]
    cleaned: list[str] = []
    previous_blank = False
    for line in lines:
        is_blank = not line
        if not (is_blank and previous_blank):
            cleaned.append(line)
        previous_blank = is_blank
    return "\n".join(cleaned).strip()


def content_hash(text: str) -> str:
    """返回规范化正文的 SHA-256；用于内容去重，不是加密文件。"""

    return hashlib.sha256(normalize_text(text).encode("utf-8")).hexdigest()


def _safe_title(filename: str) -> str:
    """从文件名得到展示标题，并移除资料集中的数字编号前缀。"""

    stem = Path(filename).stem
    stem = re.sub(r"^\d+[_\-、.\s]*", "", stem).strip()
    return stem or Path(filename).stem or "未命名文档"


def _published_at(text: str) -> str:
    """只从正文识别明确日期；没有可靠日期时返回空字符串。"""

    match = _DATE_PATTERN.search(text[:4_000])
    if not match:
        return ""
    year, month, day = (int(item) for item in match.groups())
    if month > 12 or day > 31:
        return ""
    return f"{year:04d}-{month:02d}-{day:02d}"


def _iter_docx_blocks(document: DocxDocumentType) -> Iterable[tuple[str, str, bool]]:
    """按 DOCX XML 原始顺序产出段落和表格。

    python-docx 的 document.paragraphs 与 document.tables 是两张分开的
    列表，直接拼接会打乱“段落—表格—段落”的原始顺序。遍历 body 子节点
    才能保持作者编排的上下文。
    """

    # 含 yield 的函数是生成器：调用时不会一次完成全部遍历，for 循环每索取
    # 一项才继续运行到下一个 yield，适合按原始顺序处理大文档。
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = normalize_text(paragraph.text)
            style_name = (paragraph.style.name if paragraph.style else "").lower()
            is_heading = bool(text) and (
                "heading" in style_name
                or "标题" in style_name
                or _HEADING_PATTERN.match(text) is not None
            )
            if text:
                yield "paragraph", text, is_heading
        elif tag == "tbl":
            table = Table(child, document)
            rows: list[str] = []
            for row in table.rows:
                cells = [normalize_text(cell.text).replace("\n", " ") for cell in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                yield "table", "\n".join(rows), False


def _parse_docx(data: bytes, filename: str) -> ParsedDocument:
    """解析 DOCX，并在解压前限制总展开大小以降低压缩炸弹风险。"""

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            expanded_size = sum(item.file_size for item in archive.infolist())
            if expanded_size > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise DocumentParseError("DOCX 解压后的内容超过安全限制")
        document = DocxDocument(io.BytesIO(data))
    except DocumentParseError:
        raise
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as error:
        raise DocumentParseError("DOCX 文件损坏或不是有效的 Office 文档") from error

    title = _safe_title(filename)
    current_heading = title
    section_parts: list[str] = []
    sections: list[tuple[str, str]] = []

    # 嵌套函数只服务当前 DOCX。nonlocal 表示重新绑定外层的 section_parts，
    # 而不是在 flush_section 内创建一个同名局部变量。
    def flush_section() -> None:
        nonlocal section_parts
        body = normalize_text("\n".join(section_parts))
        if body:
            sections.append((current_heading, body))
        section_parts = []

    for _, block_text, is_heading in _iter_docx_blocks(document):
        if is_heading and len(block_text) <= 120:
            flush_section()
            current_heading = block_text
        else:
            section_parts.append(block_text)
    flush_section()

    # 标题也是正文语义的一部分，必须参与内容指纹；否则两个正文相同但适用
    # 站点不同的文档可能被错误判为重复。
    full_text = normalize_text(
        "\n\n".join(f"{heading}\n{body}" for heading, body in sections)
    )
    if not full_text:
        raise DocumentParseError("DOCX 没有可提取的正文")
    return ParsedDocument(filename, title, full_text, sections, _published_at(full_text))


def _parse_pdf(data: bytes, filename: str) -> ParsedDocument:
    """按页提取 PDF 文字；扫描图片型 PDF 不做 OCR，并返回明确错误。"""

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [normalize_text(page.extract_text() or "") for page in reader.pages]
    except Exception as error:
        raise DocumentParseError("PDF 文件损坏、加密或格式无法解析") from error

    sections = [(f"第 {index} 页", text) for index, text in enumerate(pages, 1) if text]
    full_text = normalize_text("\n\n".join(text for _, text in sections))
    if not full_text:
        raise DocumentParseError("PDF 没有可提取正文，可能是扫描版；当前服务不执行 OCR")
    return ParsedDocument(
        filename, _safe_title(filename), full_text, sections, _published_at(full_text)
    )


def _parse_txt(data: bytes, filename: str) -> ParsedDocument:
    """按 UTF-8、UTF-8 BOM、GB18030 顺序尝试解码 TXT。"""

    decoded: str | None = None
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            decoded = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if decoded is None:
        raise DocumentParseError("TXT 编码无法识别，请使用 UTF-8 或 GB18030")

    text = normalize_text(decoded)
    if not text:
        raise DocumentParseError("TXT 文件为空")
    title = _safe_title(filename)
    return ParsedDocument(filename, title, text, [(title, text)], _published_at(text))


def parse_document(data: bytes, filename: str) -> ParsedDocument:
    """校验大小和扩展名，再将文件解析成统一结构。

    filename 只取 basename，既避免浏览器上传绝对路径，也防止把上级路径
    字符带入响应或 metadata。
    """

    # 这是上传和批量导入共同调用的门面。以后增加格式时在这里分派，并让
    # 新解析器仍返回 ParsedDocument，后续去重、切块代码就无需修改。
    safe_name = Path(filename).name
    suffix = Path(safe_name).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise DocumentParseError("仅支持 DOCX、PDF、TXT 文件")
    if not data:
        raise DocumentParseError("上传文件为空")
    if len(data) > MAX_UPLOAD_BYTES:
        raise DocumentParseError("单文件不能超过 25 MiB")

    if suffix == ".docx":
        return _parse_docx(data, safe_name)
    if suffix == ".pdf":
        return _parse_pdf(data, safe_name)
    return _parse_txt(data, safe_name)


def split_document(document: ParsedDocument) -> list[KnowledgeChunk]:
    """按章节切成约 800 字符、重叠 120 字符的知识块。

    每个块都带“文档标题 / 章节标题”前缀，使被单独召回的块仍有上下文。
    ID 由正文哈希和块序号决定，因此同一正文反复导入会得到相同 ID。
    """

    # 输入是格式统一后的 ParsedDocument，输出是尚未向量化的纯文本块。
    # Embedding 不放在这里，因此切块测试无需联网，也不会产生模型费用。
    digest = content_hash(document.text)
    chunks: list[KnowledgeChunk] = []
    for section_name, section_text in document.sections:
        prefix = f"文档：{document.title}\n章节：{section_name}\n"
        splitter = RecursiveCharacterTextSplitter(
            chunk_size=max(200, CHUNK_SIZE - len(prefix)),
            chunk_overlap=min(CHUNK_OVERLAP, max(0, CHUNK_SIZE - len(prefix) - 1)),
            separators=SPLIT_SEPARATORS,
            length_function=len,
        )
        for body in splitter.split_text(section_text):
            chunk_index = len(chunks)
            chunks.append(
                KnowledgeChunk(
                    chunk_id=f"{digest}:{chunk_index:05d}",
                    text=prefix + body,
                    chunk_index=chunk_index,
                    section=section_name,
                )
            )
    if not chunks:
        raise DocumentParseError("文件没有生成任何知识块")
    return chunks
