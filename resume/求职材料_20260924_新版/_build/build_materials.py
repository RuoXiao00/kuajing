from pathlib import Path
import re, json, shutil
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent.parent
PHOTO = Path(r'D:\电脑管家迁移文件\xwechat_files\wxid_szuvtcuegjoh22_8368\temp\RWTemp\2026-09\9e20f478899dc29eb19741386f9343c8\d82e296154f42c3545020b8d6cca83d0.jpg')
FONT = '微软雅黑'
GITHUB = 'https://github.com/RuoXiao00/kuajing'
LIVE = 'https://app.qiyuange.online'

def font(run, size=None, bold=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)

def setup(doc, resume=False):
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin, s.bottom_margin = Cm(1.55 if resume else 1.8), Cm(1.5 if resume else 1.7)
    s.left_margin = s.right_margin = Cm(1.75 if resume else 1.9)
    s.header_distance = s.footer_distance = Cm(.75)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Number']:
        st = doc.styles[name]
        st.font.name = FONT
        st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.size = Pt(10.5 if resume else 10.5)
        st.paragraph_format.space_after = Pt(4 if resume else 5)
        st.paragraph_format.line_spacing = Pt(17) if resume else 1.2
    for n, size in [('Title',24),('Heading 1',14),('Heading 2',11.5),('Heading 3',11)]:
        st=doc.styles[n]; st.font.size=Pt(size); st.font.bold=True
        st.paragraph_format.space_before=Pt(10 if n!='Title' else 0)
        st.paragraph_format.space_after=Pt(5)
        st.paragraph_format.keep_with_next=True
    doc.styles['Title'].paragraph_format.space_after=Pt(5)
    if resume: doc.styles['Title'].paragraph_format.line_spacing=Pt(30)
    doc.core_properties.author='黄健鸿'
    doc.core_properties.subject='个人求职材料'
    doc.core_properties.language='zh-CN'
    if not resume:
        p=s.header.paragraphs[0]
        p.add_run('黄健鸿  |  跨境阁项目面试准备'); font(p.runs[0],8)
        p=s.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        r=p.add_run('第 '); font(r,8)
        fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); p._p.append(fld)
        font(p.add_run(' 页'),8)
    # Remove template borders and Chinese page-grid snapping for compact, clean print layout.
    for element in list(doc.styles.element.iter(qn('w:pBdr'))):
        element.getparent().remove(element)
    for element in list(doc.element.iter(qn('w:pBdr'))):
        element.getparent().remove(element)
    for section in doc.sections:
        for grid in list(section._sectPr.findall(qn('w:docGrid'))):
            section._sectPr.remove(grid)
    for st in doc.styles:
        if st.type == 1:
            pp = st.element.get_or_add_pPr()
            snap = OxmlElement('w:snapToGrid'); snap.set(qn('w:val'), '0'); pp.append(snap)
    if resume:
        doc.styles['Heading 2'].paragraph_format.space_before = Pt(7)
        doc.styles['Heading 2'].paragraph_format.space_after = Pt(4)
    return doc

def ptext(doc, text, style=None, boldlead=False):
    p=doc.add_paragraph(style=style)
    if boldlead and '：' in text:
        a,b=text.split('：',1); font(p.add_run(a+'：'),bold=True); font(p.add_run(b))
    else: font(p.add_run(text))
    return p

def link(p, text, url, size=9):
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    fs=OxmlElement('w:rFonts'); fs.set(qn('w:ascii'),FONT); fs.set(qn('w:eastAsia'),FONT); pr.append(fs)
    sz=OxmlElement('w:sz'); sz.set(qn('w:val'),str(size*2)); pr.append(sz)
    color=OxmlElement('w:color'); color.set(qn('w:val'),'000000'); pr.append(color)
    r.append(pr); t=OxmlElement('w:t'); t.text=text; r.append(t); h.append(r); p._p.append(h)

def photo_anchor(p):
    # Preserve the supplied headshot without cropping or changing the person.
    shape=p.add_run().add_picture(str(PHOTO),width=Cm(2.4))
    inline=shape._inline
    cx,cy=inline.extent.cx,inline.extent.cy
    anchor=parse_xml(f'''<wp:anchor {nsdecls('wp','a','pic','r')} distT="0" distB="0" distL="114300" distR="0" simplePos="0" relativeHeight="251658240" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="0"><wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="margin"><wp:align>right</wp:align></wp:positionH><wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV><wp:extent cx="{cx}" cy="{cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapSquare wrapText="left"/></wp:anchor>''')
    for tag in ['docPr','cNvGraphicFramePr','graphic']:
        child=inline.find(qn('a:graphic') if tag=='graphic' else qn('wp:'+tag))
        if child is not None: anchor.append(child)
    inline.getparent().replace(inline,anchor)

AI_BULLETS = [
'知识库问答：用 LangGraph 拆分问题改写、召回、精排和生成节点；基于 Chroma 与 Embedding 召回 12 个候选，精排后最多选取 5 个片段，结合来源引用与缺少资料时的回复分支。',
'文档入库：按正文指纹检查重复，在持久化阶段再次校验；将解析、分块、向量写入拆为独立步骤，便于定位失败与离线测试。',
'推荐与采集：用 Playwright 采集公开商品资料，处理 ASIN 去重、价格有效性及分页异常；用 SQLite 保存每日任务和推荐快照，更新失败时保留最近可用结果。',
'图片生成：对接 Coze 工作流，支持参考图上传与可选提示词；通过请求标识避免重复提交，将生成图下载到服务端，保存失败时仅重试下载。',
'全栈交付：用 FastAPI 统一业务接口、React/Vite 实现页面与 SSE 流式问答；以 Docker Compose、Nginx HTTPS、GitHub Pages 部署，分目录持久化业务数据。',
'工程验证：编写入库去重、RAG 分支、推荐状态、图片任务及商品解析相关测试，以依赖替身隔离模型调用，覆盖异常与恢复路径。'
]
FDE_BULLETS = [
'部署交付：独立完成 React/Vite 前端到 GitHub Pages、FastAPI 后端到阿里云 ECS 的部署；配置 DNS、Nginx 反向代理和 HTTPS，后端容器端口仅绑定宿主机回环地址。',
'环境与运维：通过 Docker Compose 管理运行环境、重启策略与健康检查；为知识库、推荐、热门商品和图片配置独立持久化目录，并限制容器日志大小。',
'故障排查：处理证书验证超时、镜像拉取和依赖下载缓慢等部署问题，结合本机接口、端口监听及安全组配置分层定位；区分公网访问故障与上游采集异常。',
'业务恢复：设计推荐快照与任务状态，更新失败时保留可用结果；图片任务使用请求标识去重，已生成图片保存失败时单独重试下载，降低重复触发工作流的风险。',
'系统集成：对接模型、Embedding、Rerank 与 Coze 工作流；用 LangGraph 组织 RAG 和推荐流程，打通文档入库、流式回答、商品浏览和图片生成页面。',
'交付资料：整理 GitHub Pages 与云服务器部署说明，配套环境配置示例、检查脚本及业务测试，为环境复现、排错和后续维护提供依据。'
]

def resume(role,stem):
    doc=setup(Document(),True)
    p=doc.add_paragraph(style='Title'); font(p.add_run('黄健鸿'),24,True); photo_anchor(p)
    p.paragraph_format.right_indent=Cm(3)
    rows=[f'求职意向  {role}实习生', '广州  |  13189576720  |  ruoxiao77@gmail.com', '可随时到岗  |  每周 5 天  |  可长期实习']
    if role=='FDE 前线部署工程师': rows[-1]+='  |  接受出差与驻场'
    for i,t in enumerate(rows):
        p=ptext(doc,t); p.paragraph_format.right_indent=Cm(2.8)
        for r in p.runs: font(r,10 if i else 11, i==0)
    doc.add_heading('教育经历',2)
    ptext(doc,'仲恺农业工程学院  |  计算机科学与技术  |  本科在读')
    ptext(doc,'2024.09—2028（预计毕业）')
    doc.add_heading('个人优势',2)
    ptext(doc,('独立负责“跨境阁”全栈 AI 项目，完成知识库 RAG、商品推荐与图片工作流的开发、部署和维护。能够围绕数据、模型、接口与页面排查问题，并通过测试验证关键业务分支。' if role.startswith('AI') else '独立开发、部署并维护“跨境阁”AI 应用，具备模型服务接入、前后端联调和云服务器排错实践。能从业务流程拆解到上线运行完整推进个人项目，可接受出差与客户驻场。'))
    doc.add_heading('项目经历',2)
    ptext(doc,'跨境阁  |  全栈 AI 应用  |  独立开发与维护',boldlead=False).runs[0].bold=True
    ptext(doc,'2026.08—至今  |  个人项目')
    ptext(doc,'面向跨境选品场景，整合商品资料浏览、每日推荐、知识库问答和参考图生成。')
    for t in (AI_BULLETS if role.startswith('AI') else FDE_BULLETS):
        p=ptext(doc,t,'List Bullet',True)
        p.paragraph_format.space_after=Pt(5)
        p.paragraph_format.left_indent=Cm(.35)
        p.paragraph_format.first_line_indent=Cm(-.3)
    doc.add_heading('技术能力',2)
    skills = (['AI 应用：LangGraph、RAG、Chroma、Embedding / Rerank、Coze 工作流。','后端与数据：Python、FastAPI、SQLite、SSE、Playwright；任务状态、数据校验与接口测试。','前端与部署：React、Vite、Docker Compose、Nginx、GitHub Pages、Linux 基础排查。'] if role.startswith('AI') else ['部署与排查：Linux、Docker Compose、Nginx、HTTPS、DNS、安全组、日志与健康检查。','开发与集成：Python、FastAPI、React、SQLite、SSE；模型 API、RAG 与 Coze 工作流接入。','交付实践：环境配置、数据持久化、检查脚本、部署文档、异常恢复与业务流程验证。'])
    for t in skills: ptext(doc,t,boldlead=True)
    doc.add_heading('项目链接',2)
    p=doc.add_paragraph(); link(p,'GitHub  '+GITHUB,GITHUB,9)
    p=doc.add_paragraph(); link(p,'在线作品  '+LIVE,LIVE,9)
    doc.save(ROOT/(stem+'.docx'))

def md_inline(p, text):
    # Keep emphasis and links editable and clickable in the Word companion.
    parts=re.split(r'(\*\*.*?\*\*|\[[^\]]+\]\(https?://[^\)]+\)|`[^`]+`)',text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'): font(p.add_run(part[2:-2]),bold=True)
        elif re.match(r'^\[[^\]]+\]\(https?://',part):
            m=re.match(r'\[([^\]]+)\]\(([^\)]+)\)',part); link(p,m[1],m[2],10)
        elif part.startswith('`') and part.endswith('`'):
            r=p.add_run(part[1:-1]); font(r,10); r.font.name='Consolas'
        else: font(p.add_run(part))

def md_to_docx(path):
    doc=setup(Document())
    code=False
    for line in path.read_text(encoding='utf-8').splitlines():
        if line.startswith('```'): code=not code; continue
        if code:
            p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.05
            r=p.add_run(line or ' '); font(r,9); r.font.name='Consolas'; continue
        if not line.strip(): continue
        if line=='<!-- PAGE -->': doc.add_page_break(); continue
        if line.startswith('# '):
            p=doc.add_paragraph(line[2:],'Title')
            for r in p.runs: font(r,23,True)
        elif line.startswith('## '): doc.add_heading(line[3:],1)
        elif line.startswith('### '): doc.add_heading(line[4:],2)
        elif line.startswith('#### '): doc.add_heading(line[5:],3)
        elif line.startswith('- '): md_inline(doc.add_paragraph(style='List Bullet'),line[2:])
        else: md_inline(doc.add_paragraph(),line)
    # Keep each interview question and its answer or code exercise on one page when it fits.
    paragraphs = doc.paragraphs
    starts = [i for i,p in enumerate(paragraphs) if p.style.name == 'Heading 2']
    for i in starts:
        j = i + 1
        while j < len(paragraphs) and paragraphs[j].style.name not in ('Heading 1','Heading 2','Title'):
            j += 1
        for k in range(i,j):
            paragraphs[k].paragraph_format.keep_together = True
            paragraphs[k].paragraph_format.keep_with_next = k < j-1
    for p in paragraphs:
        if p.style.name == 'Heading 1' and p.text.startswith(('一 ', '八 模拟', '七 如何练')):
            p.paragraph_format.page_break_before = True
    doc.save(path.with_suffix('.docx'))

def write_text(name, text):
    (ROOT/name).write_text(text.strip()+'\n',encoding='utf-8')

def boss(role):
    ai=role=='AI'
    return f'''# 黄健鸿 {'AI 应用开发工程师' if ai else 'FDE 前线部署工程师'} BOSS直聘在线简历

## 基本信息与求职意向

- 姓名：黄健鸿
- 手机：13189576720
- 邮箱：ruoxiao77@gmail.com
- 当前身份：本科在读，2028 年预计毕业
- 所在城市与期望城市：广州
- 求职类型：实习
- 期望职位：{'AI 应用开发工程师 / 大模型应用开发实习生 / AI 全栈开发实习生' if ai else 'FDE 实习生 / AI 交付工程师实习生 / AI 解决方案实施实习生'}
- 到岗与出勤：可随时到岗，每周 5 天，可长期实习；具体结束时间与公司协商
- 出差与驻场：可接受
- 期望薪资：按岗位要求填写或面议，不预填具体数字

## 个人优势 直接粘贴版

{('RAG与Agent实战｜独立开发并部署全栈AI应用。仲恺农业工程学院计算机本科在读，2028年毕业。独立负责“跨境阁”，用 LangGraph、Chroma、FastAPI、React 实现知识库问答、商品推荐与图片工作流；完成 Docker、Nginx HTTPS 和 GitHub Pages 部署。项目包含流式回答、入库去重、推荐快照与图片本地保存，可提供代码和在线演示。广州求职，可随时到岗，每周5天，可长期实习。' if ai else '独立部署并维护AI项目｜可出差驻场｜广州实习。计算机本科在读，2028年毕业。独立完成“跨境阁”从模型与业务接口集成，到 Docker Compose、Nginx HTTPS 和 GitHub Pages 上线；有证书验证、网络连通、上游异常与数据持久化排查实践。能够结合日志和业务链路定位问题，整理配置示例、检查脚本及部署文档。可随时到岗，每周5天，可长期实习。')}

## 教育经历

仲恺农业工程学院｜计算机科学与技术｜本科｜2024.09—2028（预计毕业）

## 项目经历 直接粘贴版

项目名称：跨境阁

项目时间：2026.08—至今

担任角色：独立开发与维护（个人项目）

项目描述：面向跨境选品场景，整合公开商品资料浏览、每日推荐、知识库问答和参考图生成，提供可访问的前端页面和后端服务。

职责与成果：

{chr(10).join('- '+x for x in (AI_BULLETS if ai else FDE_BULLETS))}

## 技能标签

{'Python、FastAPI、RAG、LangGraph、Chroma、Embedding、Rerank、React、SSE、SQLite、Docker、Coze' if ai else 'Python、FastAPI、Docker Compose、Nginx、Linux、HTTPS、API集成、RAG、SQLite、故障排查、技术文档'}

## 作品链接

- GitHub：{GITHUB}
- 在线作品：{LIVE}

## 填写时注意

- 工作经历目前没有正式实习或工作，不将个人项目填写成公司任职；项目内容放“项目经历”。
- 当前以实习身份投递，2028 年毕业不要填成已毕业或有多年工作经验。
- 如客户端没有 FDE 职位选项，按真实岗位选择“AI 交付 / 解决方案 / 大模型应用”相关类别，具体名称以客户端为准。
- “前约25字”是文案布局策略，不是查证后的 BOSS 固定截断规则；不同页面和设备展示长度会变化。
- 平台支持在线简历及直接沟通；字段和入口以当期页面为准。参考：[BOSS功能说明](https://www.zhipin.com/web/common/protocol/feature-intro.html)、[官方简历工具](https://cv.zhipin.com/)。
'''

if __name__=='__main__':
    resume('AI 应用开发工程师','01_黄健鸿_AI应用开发工程师_简历')
    resume('FDE 前线部署工程师','02_黄健鸿_FDE前线部署工程师_简历')
    write_text('03_AI应用开发_BOSS在线简历.md',boss('AI'))
    write_text('04_FDE_BOSS在线简历.md',boss('FDE'))
    for name in ['05_AI应用开发_面试题与参考答案.md','06_FDE_面试题与参考答案.md']:
        p=ROOT/name
        if p.exists(): md_to_docx(p)
    print('生成 Word 和在线简历完成')
